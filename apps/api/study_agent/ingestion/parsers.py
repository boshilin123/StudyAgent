from collections.abc import Callable
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


def _document(content: str, **metadata: object) -> Document | None:
    normalized = "\n".join(line.rstrip() for line in content.splitlines()).strip()
    if not normalized:
        return None
    return Document(page_content=normalized, metadata=metadata)


def _parse_text(content: bytes, filename: str) -> list[Document]:
    text = content.decode("utf-8-sig")
    document = _document(text, source=filename, heading_path=[])
    return [document] if document else []


def _parse_markdown(content: bytes, filename: str) -> list[Document]:
    text = content.decode("utf-8-sig")
    documents: list[Document] = []
    headings: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        document = _document("\n".join(buffer), source=filename, heading_path=list(headings))
        if document:
            documents.append(document)
        buffer.clear()

    for line in text.splitlines():
        stripped = line.lstrip()
        if stripped.startswith("#"):
            marker = stripped.split(maxsplit=1)[0]
            if set(marker) == {"#"} and 1 <= len(marker) <= 6:
                flush()
                title = stripped[len(marker) :].strip()
                level = len(marker)
                headings[:] = headings[: level - 1]
                if title:
                    headings.append(title)
                continue
        buffer.append(line)
    flush()
    return documents


def _parse_pdf(content: bytes, filename: str) -> list[Document]:
    documents: list[Document] = []
    for page_number, page in enumerate(PdfReader(BytesIO(content)).pages, start=1):
        document = _document(
            page.extract_text() or "",
            source=filename,
            page_start=page_number,
            page_end=page_number,
            heading_path=[],
        )
        if document:
            documents.append(document)
    return documents


def _parse_docx(content: bytes, filename: str) -> list[Document]:
    documents: list[Document] = []
    headings: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        document = _document("\n".join(buffer), source=filename, heading_path=list(headings))
        if document:
            documents.append(document)
        buffer.clear()

    for paragraph in DocxDocument(BytesIO(content)).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            flush()
            try:
                level = int(style_name.split()[-1])
            except ValueError:
                level = 1
            headings[:] = headings[: max(0, level - 1)]
            headings.append(text)
        else:
            buffer.append(text)
    flush()
    return documents


def _parse_pptx(content: bytes, filename: str) -> list[Document]:
    documents: list[Document] = []
    for slide_number, slide in enumerate(Presentation(BytesIO(content)).slides, start=1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        title = texts[0] if texts else f"第 {slide_number} 页"
        document = _document(
            "\n".join(texts),
            source=filename,
            page_start=slide_number,
            page_end=slide_number,
            heading_path=[title],
        )
        if document:
            documents.append(document)
    return documents


def _parse_xlsx(content: bytes, filename: str) -> list[Document]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    documents: list[Document] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    rows.append("\t".join(values).rstrip())
            document = _document("\n".join(rows), source=filename, heading_path=[worksheet.title])
            if document:
                documents.append(document)
    finally:
        workbook.close()
    return documents


PARSERS: dict[str, Callable[[bytes, str], list[Document]]] = {
    ".txt": _parse_text,
    ".md": _parse_markdown,
    ".markdown": _parse_markdown,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".pptx": _parse_pptx,
    ".xlsx": _parse_xlsx,
}


def parse_document(content: bytes, filename: str) -> list[Document]:
    suffix = Path(filename).suffix.lower()
    parser = PARSERS.get(suffix)
    if parser is None:
        raise ValueError(f"unsupported parser for suffix: {suffix}")
    documents = parser(content, filename)
    if not documents:
        raise ValueError("document contains no extractable text")
    return documents
