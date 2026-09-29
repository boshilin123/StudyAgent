from io import BytesIO
from uuid import uuid4

from docx import Document as DocxDocument
from langchain_openai import OpenAIEmbeddings
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas

from study_agent.config import Settings
from study_agent.ingestion.chunking import build_chunks
from study_agent.ingestion.embeddings import (
    DeterministicEmbeddings,
    TeiEmbeddings,
    build_embeddings,
)
from study_agent.ingestion.parsers import parse_document


def test_markdown_parser_preserves_heading_path_and_chunks() -> None:
    documents = parse_document(
        "# 网络\nTCP 提供可靠传输。\n\n## 握手\n客户端先发送 SYN。".encode(),
        "网络.md",
    )
    chunks = build_chunks(
        documents,
        material_id=uuid4(),
        chunk_size=30,
        chunk_overlap=5,
    )

    assert len(documents) == 2
    assert documents[1].metadata["heading_path"] == ["网络", "握手"]
    assert chunks
    assert all(chunk.content_hash and chunk.token_count > 0 for chunk in chunks)


def test_deterministic_embeddings_are_stable() -> None:
    embeddings = DeterministicEmbeddings(32)
    first = embeddings.embed_query("TCP 三次握手")
    second = embeddings.embed_query("TCP 三次握手")

    assert first == second
    assert len(first) == 32


def test_tei_embedding_provider_can_run_without_api_key() -> None:
    embeddings, model_name = build_embeddings(
        Settings(
            _env_file=None,
            embedding_provider="tei",
            embedding_base_url="http://embedding:80",
            embedding_model="BAAI/bge-large-zh-v1.5",
        )
    )

    assert isinstance(embeddings, TeiEmbeddings)
    assert model_name == "BAAI/bge-large-zh-v1.5"


def test_openai_compatible_provider_sends_string_inputs() -> None:
    embeddings, model_name = build_embeddings(
        Settings(
            _env_file=None,
            embedding_provider="openai_compatible",
            embedding_base_url="https://example.com/compatible-mode/v1",
            embedding_api_key="test-key",
            embedding_model="text-embedding-v4",
            embedding_dimensions=1024,
        )
    )

    assert isinstance(embeddings, OpenAIEmbeddings)
    assert embeddings.check_embedding_ctx_length is False
    assert embeddings.chunk_size == 10
    assert model_name == "text-embedding-v4"


def test_pdf_docx_pptx_and_xlsx_parsers_extract_text() -> None:
    pdf_buffer = BytesIO()
    canvas = Canvas(pdf_buffer)
    canvas.drawString(72, 720, "TCP reliable transport")
    canvas.save()

    docx_buffer = BytesIO()
    docx = DocxDocument()
    docx.add_heading("TCP", level=1)
    docx.add_paragraph("Reliable transport")
    docx.save(docx_buffer)

    pptx_buffer = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "TCP"
    slide.placeholders[1].text = "Three-way handshake"
    presentation.save(pptx_buffer)

    xlsx_buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Protocols"
    sheet.append(["name", "property"])
    sheet.append(["TCP", "reliable"])
    workbook.save(xlsx_buffer)

    samples = [
        (pdf_buffer.getvalue(), "sample.pdf", "TCP"),
        (docx_buffer.getvalue(), "sample.docx", "Reliable"),
        (pptx_buffer.getvalue(), "sample.pptx", "handshake"),
        (xlsx_buffer.getvalue(), "sample.xlsx", "reliable"),
    ]
    for content, filename, expected in samples:
        documents = parse_document(content, filename)
        assert expected in "\n".join(document.page_content for document in documents)
